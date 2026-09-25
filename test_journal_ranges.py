"""The journal's autofilter / sort ranges must follow appended rows.

Excel stores both as fixed A1-style strings, so appending trades underneath
leaves them behind. By August 2026 the filter on Trades.xlsx still ended at row
483 while the data ran to 504 — the whole month sat outside the filter and read
as missing in Excel even though every row was in the file.

The module under test has a hyphenated filename and so can't be imported with a
plain `import`; loaded here via importlib, which is also why the existing
test_trade_formatter.py doesn't run.

Run: python3 -m unittest test_journal_ranges -v
"""

import importlib.util
import unittest
from pathlib import Path

import openpyxl

_SRC = Path(__file__).with_name("trade-log-formatter.py")
_spec = importlib.util.spec_from_file_location("trade_log_formatter", _SRC)
tlf = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tlf)

extend = tlf._extend_journal_table_ranges


def _sheet(filter_ref="A1:AI483", sort_ref="A2:AI493", cond_ref="P1:P493"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trades"
    ws["A1"] = "Symbol"
    if filter_ref:
        ws.auto_filter.ref = filter_ref
        if sort_ref:
            ws.auto_filter.add_sort_condition(cond_ref, descending=True)
            ws.auto_filter.sortState.ref = sort_ref
    return ws


class ExtendRangesTests(unittest.TestCase):
    def test_filter_and_sort_follow_the_new_last_row(self):
        ws = _sheet()
        extend(ws, 504)
        self.assertEqual(ws.auto_filter.ref, "A1:AI504")
        self.assertEqual(ws.auto_filter.sortState.ref, "A2:AI504")
        self.assertEqual(ws.auto_filter.sortState.sortCondition[0].ref, "P2:P504")

    def test_sort_range_excludes_the_header_row(self):
        # The live workbook had P1:P493 — sorting with the header inside the
        # range drags the header into the data.
        ws = _sheet(cond_ref="P1:P493")
        extend(ws, 504)
        self.assertTrue(ws.auto_filter.sortState.sortCondition[0].ref.startswith("P2:"))

    def test_the_users_chosen_width_is_preserved(self):
        # Widening to ws.max_column would silently pull unrelated columns into
        # the filter; keep whatever span they set.
        ws = _sheet(filter_ref="A1:M483", sort_ref="A2:M493", cond_ref="F1:F493")
        extend(ws, 504)
        self.assertEqual(ws.auto_filter.ref, "A1:M504")
        self.assertEqual(ws.auto_filter.sortState.ref, "A2:M504")
        self.assertEqual(ws.auto_filter.sortState.sortCondition[0].ref, "F2:F504")

    def test_a_sheet_with_no_filter_is_left_alone(self):
        ws = _sheet(filter_ref=None)
        extend(ws, 504)
        self.assertIsNone(ws.auto_filter.ref)

    def test_a_filter_with_no_saved_sort_still_extends(self):
        ws = _sheet(sort_ref=None)
        extend(ws, 504)
        self.assertEqual(ws.auto_filter.ref, "A1:AI504")

    def test_an_empty_sheet_is_a_no_op(self):
        ws = _sheet()
        extend(ws, 1)
        self.assertEqual(ws.auto_filter.ref, "A1:AI483")

    def test_it_is_idempotent(self):
        ws = _sheet()
        extend(ws, 504)
        extend(ws, 504)
        self.assertEqual(ws.auto_filter.ref, "A1:AI504")
        self.assertEqual(ws.auto_filter.sortState.sortCondition[0].ref, "P2:P504")

    def test_it_shrinks_too_if_rows_are_ever_removed(self):
        # Not a scenario the formatter creates, but the range should describe
        # the data rather than only ever growing.
        ws = _sheet(filter_ref="A1:AI504", sort_ref="A2:AI504", cond_ref="P2:P504")
        extend(ws, 300)
        self.assertEqual(ws.auto_filter.ref, "A1:AI300")


class PnlFormulaTests(unittest.TestCase):
    """The sheet's carried-down P/L formula is long-biased.

    `=(ExitQty*ExitPrice)-(EntryPrice*Qty)` inverts the sign on a short: BMNR
    (45 short at 49.13, covered at 49.608) read +21.51 in Excel against a real
    -21.51. QuantForge already flipped the sign for shorts when recomputing,
    so the app and the sheet disagreed on exactly that row.
    """

    HEADERS = {"Symbol": 1, "Qty": 2, "Side": 3, "Entry Price": 4,
               "Exit Qty": 10, "Exit Price": 11}

    @staticmethod
    def _eval(side, qty, entry, exit_qty, exit_price):
        """The arithmetic the authored formula encodes, in Python."""
        proceeds, cost = exit_qty * exit_price, entry * qty
        return (cost - proceeds) if side.upper().startswith("SHORT") else (proceeds - cost)

    def test_it_branches_on_side(self):
        f = tlf._pnl_formula(self.HEADERS, 504)
        self.assertIn('LEFT(UPPER(C504),5)="SHORT"', f)
        self.assertIn("(D504*B504)-(J504*K504)", f)   # short: cost - proceeds
        self.assertIn("(J504*K504)-(D504*B504)", f)   # long:  proceeds - cost

    def test_a_losing_short_is_a_loss(self):
        self.assertAlmostEqual(self._eval("SHORT", 45, 49.13, 45, 49.608), -21.51, places=2)

    def test_a_winning_short_is_a_gain(self):
        self.assertAlmostEqual(self._eval("SHORT", 45, 49.608, 45, 49.13), 21.51, places=2)

    def test_longs_are_unchanged_by_the_new_formula(self):
        self.assertAlmostEqual(self._eval("LONG", 20, 51.6625, 20, 51.72), 1.15, places=2)

    def test_it_books_only_the_shares_actually_sold(self):
        # Matches the app's (exit_px * sold) - (entry * qty).
        self.assertAlmostEqual(self._eval("LONG", 35, 10.2, 15, 12.155714285714286),
                               -174.66, places=2)

    def test_it_declines_when_a_column_is_missing(self):
        # Caller then falls back to copying the previous row's formula rather
        # than writing a broken reference.
        self.assertIsNone(tlf._pnl_formula({"Qty": 2, "Side": 3}, 5))

    def test_it_follows_the_sheets_actual_column_positions(self):
        moved = {"Symbol": 1, "Qty": 5, "Side": 7, "Entry Price": 9,
                 "Exit Qty": 12, "Exit Price": 14}
        f = tlf._pnl_formula(moved, 10)
        self.assertIn("G10", f)          # Side
        self.assertIn("(L10*N10)", f)    # Exit Qty * Exit Price
        self.assertIn("(I10*E10)", f)    # Entry Price * Qty


if __name__ == "__main__":
    unittest.main()
